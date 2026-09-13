"""把受控的 PDF、文本和 Office 文件解析成 Markdown 派生文件。"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Protocol

from pypdf import PdfReader


class DocumentParser(Protocol):
    supported_extensions: frozenset[str]

    def parse(self, source: Path, output_dir: Path, *, max_chars: int) -> "ParsedDocument":
        ...


@dataclass(frozen=True)
class ParsedDocument:
    character_count: int
    files: tuple[str, ...]


class DocumentParseError(ValueError):
    """解析失败，错误码不会包含第三方库原始路径或堆栈。"""

    def __init__(self, message: str, code: str = "attachment_parse_failed"):
        super().__init__(message)
        self.error_code = code


def _plain(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def _write_parts(output_dir: Path, title: str, text: str, max_chars: int) -> ParsedDocument:
    text = text.strip()
    if len(text) > max_chars:
        raise DocumentParseError("解析后的附件内容过大。", "parse_output_too_large")
    output_dir.mkdir(parents=True, exist_ok=True)
    part_size = 100_000
    parts: list[str] = []
    if text:
        for offset in range(0, len(text), part_size):
            filename = f"part-{len(parts) + 1:03d}.md"
            (output_dir / filename).write_text(text[offset : offset + part_size], encoding="utf-8")
            parts.append(filename)
    index_lines = [f"# {title}", "", f"字符数：{len(text)}", ""]
    index_lines.extend(f"- [{filename}]({filename})" for filename in parts)
    (output_dir / "index.md").write_text("\n".join(index_lines) + "\n", encoding="utf-8")
    return ParsedDocument(
        character_count=len(text),
        files=tuple(["index.md", *parts]),
    )


def _parse_text(source: Path) -> str:
    return source.read_text(encoding="utf-8-sig")


def _parse_csv(source: Path) -> str:
    rows: list[list[str]] = []
    with source.open(encoding="utf-8-sig", newline="") as handle:
        rows = [[_plain(cell) for cell in row] for row in csv.reader(handle)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(normalized[0]) + " |", "| " + " | ".join("---" for _ in range(width)) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _parse_json(source: Path) -> str:
    try:
        return "```json\n" + json.dumps(json.loads(_parse_text(source)), ensure_ascii=False, indent=2) + "\n```"
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise DocumentParseError("JSON 文件格式无效。", "attachment_invalid") from exc


def _parse_pdf(source: Path) -> str:
    try:
        reader = PdfReader(str(source), strict=False)
        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if len(text) < 2:
                raise DocumentParseError("PDF 包含没有文本层的页面。", "pdf_no_text_layer")
            pages.append(f"## 第 {index} 页\n\n{text}")
        return "\n\n".join(pages)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("PDF 解析失败。") from exc


def _parse_docx(source: Path) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(source.read_bytes()))
        sections = [_plain(paragraph.text) for paragraph in document.paragraphs if _plain(paragraph.text)]
        for table in document.tables:
            sections.append("\n".join("| " + " | ".join(_plain(cell.text) for cell in row.cells) + " |" for row in table.rows))
        return "\n\n".join(sections)
    except Exception as exc:
        raise DocumentParseError("DOCX 解析失败。") from exc


def _parse_xlsx(source: Path) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(source.read_bytes()),
            read_only=True,
            data_only=True,
        )
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows = [[_plain(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized = [row + [""] * (width - len(row)) for row in rows]
            lines = [f"## {sheet.title}", "", "| " + " | ".join(normalized[0]) + " |", "| " + " | ".join("---" for _ in range(width)) + " |"]
            lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
            sections.append("\n".join(lines))
        workbook.close()
        return "\n\n".join(sections)
    except Exception as exc:
        raise DocumentParseError("XLSX 解析失败。") from exc


def _parse_pptx(source: Path) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(BytesIO(source.read_bytes()))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and _plain(shape.text):
                    texts.append(_plain(shape.text))
                if getattr(shape, "has_table", False):
                    texts.extend(
                        "| " + " | ".join(_plain(cell.text) for cell in row.cells) + " |"
                        for row in shape.table.rows
                    )
            if texts:
                slides.append(f"## 第 {index} 页\n\n" + "\n\n".join(texts))
        return "\n\n".join(slides)
    except Exception as exc:
        raise DocumentParseError("PPTX 解析失败。") from exc


def parse_document(
    source: Path,
    output_dir: Path,
    *,
    max_chars: int,
    original_name: str | None = None,
) -> ParsedDocument:
    """按上传时的安全展示名选择解析器，物理原文名可以固定为 ``blob``。"""

    display_name = original_name or source.name
    extension = Path(display_name).suffix.lower()
    try:
        if extension in {".txt", ".md"}:
            text = _parse_text(source)
        elif extension == ".csv":
            text = _parse_csv(source)
        elif extension == ".json":
            text = _parse_json(source)
        elif extension == ".pdf":
            text = _parse_pdf(source)
        elif extension == ".docx":
            text = _parse_docx(source)
        elif extension == ".xlsx":
            text = _parse_xlsx(source)
        elif extension == ".pptx":
            text = _parse_pptx(source)
        else:
            raise DocumentParseError("暂不支持该附件类型。", "unsupported_attachment_type")
    except DocumentParseError:
        raise
    except (OSError, UnicodeError) as exc:
        raise DocumentParseError("附件读取失败。") from exc
    return _write_parts(output_dir, display_name, text, max_chars)
